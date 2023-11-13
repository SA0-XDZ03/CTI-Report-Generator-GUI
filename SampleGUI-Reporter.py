import tkinter as tk
from tkinter import filedialog
from tkinter import scrolledtext
import docx

def open_template():
    template_path = filedialog.askopenfilename(filetypes=[("DOCX files", "*.docx")])
    template_selector.set(template_path)

def generate_report():
    template_path = template_selector.get()
    save_path = filedialog.asksaveasfilename(defaultextension=".docx")

    if template_path and save_path:
        doc = docx.Document(template_path)
        for placeholder, entry in placeholders.items():
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    text = run.text
                    text = text.replace(placeholder, entry.get())
                    run.text = text
        doc.save(save_path)
        status_label.config(text=f"Report saved to {save_path}")
    else:
        status_label.config(text="Please select a template and save location.")


# Create the main window
root = tk.Tk()
root.title("Auto Document Report Generator")

# Create a template selector
template_selector = tk.StringVar()
template_selector.set("")
template_selector_label = tk.Label(root, text="Select a DOCX Template:")
template_selector_label.pack()
template_selector_entry = tk.Entry(root, textvariable=template_selector, state="readonly")
template_selector_entry.pack()
template_selector_button = tk.Button(root, text="Browse", command=open_template)
template_selector_button.pack()

# Create input fields
placeholders = {
    "<header_title>": "Header Title",
    "<document_number>": "12345",
    "<report_number>": "RPT123",
    "<report_type>": "Type A",
    "<report_severity>": "High",
    "<summary>": "This is the executive summary.",
    "<details_tech>": "Technical details go here.",
    "<cve_details>": "CVE-2022-12345, CVE-2022-54321",
    "<malware_details>": "Malware details here.",
    "<ioc_details>": "IOC details here.",
    "<images>": "Image paths go here",
    "<date_time>": "2023-10-13",
    "<author>": "John Doe",
    "<references>": "https://example.com/reference",
    "<footer_note>": "Footer note",
    "<footer_stamp>": "Footer stamp"
}

for placeholder, default_value in placeholders.items():
    label = tk.Label(root, text=placeholder)
    label.pack()
    entry = tk.Entry(root, width=40)
    entry.insert(0, default_value)
    placeholders[placeholder] = entry
    entry.pack()

# Create a save button
save_button = tk.Button(root, text="Generate Report", command=generate_report)
save_button.pack()

# Create a status label
status_label = tk.Label(root, text="")
status_label.pack()

# Run the GUI
root.mainloop()
